# -*- coding: utf-8 -*-
"""
⓪ 준비 상태 점검
- 폴더/CSV/DICOM/엑셀 양식/MicroDicom 설치 위치를 확인해서 보고서를 만듭니다.
- 아무 파일도 수정하지 않습니다(읽기만 함). 보고서만 labeling\setup_report.txt 로 저장.

실행: step0_check.bat 더블클릭  (또는 명령창에서  py check_setup.py)
"""
from __future__ import annotations

import csv
import platform
import sys
from collections import Counter
from pathlib import Path

import config

lines: list[str] = []


def out(text: str = "") -> None:
    print(text)
    lines.append(text)


def section(title: str) -> None:
    out()
    out("=" * 70)
    out(title)
    out("=" * 70)


def ok(flag: bool) -> str:
    return "[OK]  " if flag else "[없음]"


# ── 1. 파이썬/패키지 ───────────────────────────────────────
def check_python() -> None:
    section("1. 파이썬과 패키지")
    out(f"Python {sys.version.split()[0]}  ({platform.platform()})")
    for name in ["numpy", "pydicom", "PIL", "openpyxl"]:
        try:
            mod = __import__(name)
            out(f"[OK]   {name} {getattr(mod, '__version__', '')}")
        except Exception as e:
            out(f"[없음] {name}  → install.bat 을 먼저 실행하세요 ({e})")
    try:
        import tkinter
        out(f"[OK]   tkinter {tkinter.TkVersion}")
    except Exception as e:
        out(f"[없음] tkinter ({e}) → python.org 설치본에는 기본 포함됩니다")


# ── 2. 폴더/파일 ──────────────────────────────────────────
def check_paths() -> None:
    section("2. 폴더와 파일")
    out(f"{ok(config.PROJECT_DIR.is_dir())} 프로젝트 폴더: {config.PROJECT_DIR}")
    if not config.PROJECT_DIR.is_dir():
        out("       → 폴더 이름/위치가 다르면 tool\\config.py 의 PROJECT_DIR 을 고치세요.")
    inside = config.PROJECT_DIR in config.TOOL_DIR.parents
    out(f"{'[OK]  ' if inside else '[참고]'} 코드 위치: {config.TOOL_DIR}"
        + ("" if inside else "  (프로젝트 폴더 밖이지만 결과물은 프로젝트 폴더에 저장됩니다. "
                             "관리하기 쉽게 연구멘토링\\tool 로 옮기는 것을 권장)"))
    for label, p in [("data 폴더", config.DATA_DIR),
                     ("train_images", config.IMAGES_DIR),
                     ("series CSV", config.SERIES_CSV),
                     ("좌표 CSV", config.COORDS_CSV),
                     ("labeling 폴더", config.LABELING_DIR),
                     ("엑셀 양식", config.EXCEL_PATH),
                     ("captures 폴더", config.CAPTURE_DIR)]:
        out(f"{ok(p.exists())} {label}: {p}")
    if config.PROJECT_DIR.is_dir():
        show_tree(config.PROJECT_DIR)
        if not (config.SERIES_CSV.exists() and config.EXCEL_PATH.exists()):
            find_misplaced(config.PROJECT_DIR)


def _children(folder: Path) -> list[Path]:
    try:
        return sorted(folder.iterdir(), key=lambda p: (not p.is_dir(), p.name.lower()))
    except OSError:
        return []


def show_tree(root: Path, max_depth: int = 3, max_items: int = 15) -> None:
    """프로젝트 폴더 안 구조 (train_images 속 study 폴더는 개수만 표시)."""
    out()
    out("■ 프로젝트 폴더 안 실제 구조")

    def walk(folder: Path, depth: int) -> None:
        kids = [k for k in _children(folder) if k.name != "__pycache__"]
        if folder.name.lower() == "train_images":
            out(f"{'    ' * depth}(study 폴더 {sum(k.is_dir() for k in kids)}개)")
            return
        for k in kids[:max_items]:
            out(f"{'    ' * depth}{k.name}{chr(92) if k.is_dir() else ''}")
            if k.is_dir() and depth + 1 < max_depth:
                walk(k, depth + 1)
        if len(kids) > max_items:
            out(f"{'    ' * depth}... 외 {len(kids) - max_items}개")

    walk(root, 1)


def find_misplaced(root: Path, max_depth: int = 5) -> None:
    """CSV·엑셀·train_images 가 예상과 다른 곳에 있으면 찾아서 알려 줌."""
    targets = {config.SERIES_CSV.name, config.COORDS_CSV.name}
    found = []

    def walk(folder: Path, depth: int) -> None:
        for k in _children(folder):
            if k.is_dir():
                if k.name.lower() == "train_images":
                    found.append(("train_images 폴더", k))
                elif depth < max_depth and k.name != "__pycache__":
                    walk(k, depth + 1)
            elif k.name in targets or (k.suffix.lower() == ".xlsx" and not k.name.startswith("~$")):
                found.append(("파일", k))

    walk(root, 1)
    out()
    out("■ 프로젝트 폴더 전체에서 찾은 데이터/엑셀 위치")
    for kind, p in found:
        out(f"  {kind}: {p}")
    if not found:
        out("  CSV, train_images, 엑셀(.xlsx)을 하나도 찾지 못했습니다. 파일을 이 폴더 안으로 옮겨 주세요.")


# ── 3. CSV ────────────────────────────────────────────────
def check_csvs() -> None:
    section("3. CSV 내용")
    for p in [config.SERIES_CSV, config.COORDS_CSV]:
        if not p.exists():
            out(f"[없음] {p.name}")
            continue
        with open(p, newline="", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            rows = list(reader)
        out(f"■ {p.name}: {len(rows)}행, 열 = {reader.fieldnames}")
        if rows:
            out(f"  첫 행 예시: {dict(rows[0])}")
        for col in ["series_description", "condition", "level"]:
            if reader.fieldnames and col in reader.fieldnames:
                counts = Counter(r[col] for r in rows)
                out(f"  {col} 종류: {dict(counts)}")


# ── 4. 다운로드 상태 ──────────────────────────────────────
def check_downloads() -> str | None:
    """다운로드 현황을 출력하고, DICOM 샘플로 쓸 Sagittal T2 폴더 하나를 돌려줌."""
    section("4. 다운로드 상태 (train_images)")
    if not (config.IMAGES_DIR.is_dir() and config.SERIES_CSV.exists()):
        out("train_images 폴더 또는 series CSV 가 없어 건너뜁니다.")
        return None
    import dataset

    series_map = dataset.read_series_descriptions()
    downloaded = dataset.downloaded_study_ids()
    out(f"CSV에 있는 study 수: {len(series_map)}")
    out(f"train_images 안 study 폴더 수: {len(downloaded)}")

    n_t2 = n_t1 = n_both = 0
    not_in_csv, empty_series, gap_series = [], [], []
    sample_dir = None
    for study in downloaded:
        desc_of = dict(series_map.get(study, []))
        has = {"t2": False, "t1": False}
        for sdir in (config.IMAGES_DIR / study).iterdir():
            if not sdir.is_dir():
                continue
            files = dataset.list_dicom_files(sdir)
            if sdir.name not in desc_of:
                not_in_csv.append(f"{study}/{sdir.name}")
            if not files:
                empty_series.append(f"{study}/{sdir.name}")
                continue
            nums = [int(f.stem) for f in files if f.stem.isdigit()]
            if nums and max(nums) - min(nums) + 1 != len(nums):
                gap_series.append(f"{study}/{sdir.name} ({len(nums)}장, {min(nums)}~{max(nums)})")
            desc = desc_of.get(sdir.name, "")
            if desc == config.DESC_SAG_T2:
                has["t2"] = True
                sample_dir = sample_dir or sdir
            elif desc == config.DESC_SAG_T1:
                has["t1"] = True
        n_t2 += has["t2"]
        n_t1 += has["t1"]
        n_both += has["t2"] and has["t1"]

    out(f"Sagittal T2/STIR 가 있는 study: {n_t2}")
    out(f"Sagittal T1 이 있는 study:      {n_t1}")
    out(f"둘 다 있는 study:               {n_both}")
    out(f"CSV에 없는 series 폴더: {len(not_in_csv)} {not_in_csv[:5]}")
    out(f"비어 있는 series 폴더: {len(empty_series)} {empty_series[:5]}")
    out(f"instance 번호가 중간에 빠진 series(다운로드 중단 의심): {len(gap_series)} {gap_series[:5]}")
    return str(sample_dir) if sample_dir else None


# ── 5. DICOM 샘플 ─────────────────────────────────────────
def check_dicom(sample_dir: str | None) -> None:
    section("5. DICOM 샘플 (Sagittal T2 한 장)")
    if not sample_dir:
        out("샘플로 읽을 Sagittal T2 폴더가 없습니다.")
        return
    try:
        import pydicom
        import dataset

        files = dataset.list_dicom_files(Path(sample_dir))
        f = files[len(files) // 2]
        ds = pydicom.dcmread(f)
        out(f"파일: {f}")
        out(f"시리즈 장수: {len(files)}")
        for tag in ["InstanceNumber", "Rows", "Columns", "PixelSpacing", "SliceThickness",
                    "BitsStored", "PhotometricInterpretation", "WindowCenter", "WindowWidth",
                    "RescaleSlope", "RescaleIntercept", "ImagePositionPatient"]:
            out(f"  {tag}: {getattr(ds, tag, '(없음)')}")
        ts = ds.file_meta.TransferSyntaxUID
        out(f"  TransferSyntax: {ts} ({ts.name}), 압축={ts.is_compressed}")
        out(f"  파일명 숫자와 InstanceNumber 일치: {f.stem == str(getattr(ds, 'InstanceNumber', ''))}")
        arr = ds.pixel_array
        out(f"  픽셀 읽기 [OK] shape={arr.shape}, dtype={arr.dtype}, min={arr.min()}, max={arr.max()}")
    except Exception as e:
        out(f"[문제] DICOM 읽기 실패: {e!r}")


# ── 6. 엑셀 양식 ──────────────────────────────────────────
def check_excel() -> None:
    section("6. 엑셀 양식")
    p = config.EXCEL_PATH
    if not p.exists():
        out(f"[없음] {p}")
        return
    import excel_io
    from openpyxl.utils import get_column_letter

    if excel_io.lock_file_exists(p):
        out("[주의] 지금 Excel에서 이 파일이 열려 있습니다. 라벨링할 때는 닫아 두세요.")
    wb, warns = excel_io.load(p)
    theme = excel_io.theme_colors(wb)
    out(f"파일: {p.name} ({p.stat().st_size:,} bytes)")
    out(f"openpyxl 경고: {warns if warns else '없음'}")
    out(f"시트 목록: {wb.sheetnames}")
    try:
        names = list(wb.defined_names.keys())
    except Exception:
        names = []
    out(f"이름 정의: {names if names else '없음'}")

    for ws in wb.worksheets:
        out()
        out(f"── 시트 [{ws.title}]  크기={ws.dimensions}, 최대행={ws.max_row}, 최대열={ws.max_column}, "
            f"틀고정={ws.freeze_panes}, 숨김={ws.sheet_state}")

        if excel_io.norm(ws.title) == "readme":
            out("  (README 내용 전체)")
            for row in ws.iter_rows():
                for c in row:
                    if c.value not in (None, ""):
                        out(f"  {c.coordinate}: {c.value}")
            continue

        hr = excel_io.find_header_row(ws) or 1
        out(f"  머리글 행: {hr}")
        if ws.tables:
            out(f"  표(Table): {[(t.name, t.ref) for t in ws.tables.values()]}")
        if ws.merged_cells.ranges:
            out(f"  병합 셀: {[str(r) for r in list(ws.merged_cells.ranges)[:10]]}")
        for c in excel_io.header_columns(ws, hr, theme):
            fg = ws.cell(hr, c["col"]).fill.fgColor
            if fg.type == "theme":
                raw = f"theme={fg.theme}, tint={fg.tint:.2f}"
            else:
                raw = f"{fg.type}={getattr(fg, fg.type, '')}"
            n_formula = sum(1 for r in range(hr + 1, ws.max_row + 1)
                            if excel_io.is_formula(ws.cell(r, c["col"]).value))
            n_value = sum(1 for r in range(hr + 1, ws.max_row + 1)
                          if ws.cell(r, c["col"]).value not in (None, "")
                          and not excel_io.is_formula(ws.cell(r, c["col"]).value))
            options = excel_io.list_validation_options(ws, c["col"], hr + 1)
            out(f"  {c['letter']:>3} | {c['text']:<28} | {excel_io.KIND_KO[c['kind']]:<10} "
                f"#{c['hex']} ({raw}) | 수식칸 {n_formula}, 값칸 {n_value}"
                + (f" | 드롭다운 {options}" if options else ""))
        # 머리글 아래 3행 미리보기 (수식은 그대로 표시)
        for r in range(hr + 1, min(hr + 4, ws.max_row) + 1):
            vals = [f"{get_column_letter(c.column)}={c.value!r}" for c in ws[r] if c.value not in (None, "")]
            out(f"  {r}행: {vals if vals else '(비어 있음)'}")
        for dv in ws.data_validations.dataValidation:
            out(f"  데이터유효성: type={dv.type}, 범위={dv.sqref}, 값={dv.formula1}")
        for rng in ws.conditional_formatting:
            out(f"  조건부서식: {rng.sqref} ({len(rng.rules)}개 규칙)")


# ── 7. MicroDicom ────────────────────────────────────────
def check_microdicom() -> None:
    section("7. MicroDicom 설치 위치 (⑥단계용)")
    found = []
    for base in [Path(r"C:\Program Files"), Path(r"C:\Program Files (x86)")]:
        if not base.is_dir():
            continue
        for d in base.iterdir():
            if "microdicom" in d.name.lower() and d.is_dir():
                found += [str(x) for x in d.glob("*.exe")]
    out("\n".join(found) if found else "Program Files 안에서 찾지 못했습니다 (⑥단계에서 직접 지정 가능).")


def main() -> None:
    out("LumbarDISC 쉬모를 결절 라벨링 도구 — 준비 상태 점검")
    for step in (check_python, check_paths, check_csvs):
        try:
            step()
        except Exception as e:
            out(f"[오류] {step.__name__}: {e!r}")
    try:
        sample = check_downloads()
    except Exception as e:
        out(f"[오류] check_downloads: {e!r}")
        sample = None
    for step in (lambda: check_dicom(sample), check_excel, check_microdicom):
        try:
            step()
        except Exception as e:
            out(f"[오류] {e!r}")

    if not config.PROJECT_DIR.is_dir():
        print(f"\n프로젝트 폴더가 없어 보고서를 저장하지 않았습니다: {config.PROJECT_DIR}")
        return
    report = config.LABELING_DIR / "setup_report.txt"
    try:
        report.parent.mkdir(exist_ok=True)
        report.write_text("\n".join(lines), encoding="utf-8")
        print(f"\n보고서 저장: {report}")
    except Exception as e:
        print(f"\n보고서 저장 실패: {e}")


if __name__ == "__main__":
    main()
