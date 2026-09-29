# -*- coding: utf-8 -*-
"""
엑셀 양식을 다루는 공통 함수.

핵심 규칙
- 머리글(헤더) 칸의 배경색으로 열의 성격을 판별합니다.
    회색 = 프로그램이 입력하는 열 / 초록색 = 수식 / 노란색 = 사람이 입력
- 프로그램은 '회색' 열에만 값을 씁니다. 수식이 들어 있는 칸은 절대 덮어쓰지 않습니다.
- 저장은 임시 파일에 먼저 쓴 뒤 교체합니다. 엑셀이 파일을 열고 있으면
  교체가 실패하므로 원본은 그대로 남고 ExcelLockedError 가 납니다.
"""
from __future__ import annotations

import colorsys
import os
import re
import shutil
import time
import warnings
import xml.etree.ElementTree as ET
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.utils import get_column_letter, range_boundaries

import config


class ExcelLockedError(Exception):
    """엑셀 파일이 다른 프로그램(주로 Excel)에서 열려 있어 저장하지 못함."""


# ── 머리글 글자 정규화 ─────────────────────────────────────
def norm(text) -> str:
    """'Sagittal T2 series_id' → 'sagittalt2seriesid' (대소문자·공백·기호 무시)"""
    if text is None:
        return ""
    return re.sub(r"[^0-9a-z가-힣]", "", str(text).lower())


def is_formula(value) -> bool:
    if value is None:
        return False
    if isinstance(value, str):
        return value.startswith("=")
    # ArrayFormula, DataTableFormula 등 수식 객체
    return value.__class__.__name__.endswith("Formula")


# ── 색 판별 ───────────────────────────────────────────────
# 엑셀 테마 색 번호 → 테마 XML 안 이름 (앞 4개는 순서가 뒤바뀌는 엑셀 규칙)
_THEME_ORDER = ["lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3",
                "accent4", "accent5", "accent6", "hlink", "folHlink"]
_DEFAULT_THEME = {  # Office 기본 테마 (테마 정보를 못 읽을 때 사용)
    "lt1": "FFFFFF", "dk1": "000000", "lt2": "E7E6E6", "dk2": "44546A",
    "accent1": "4472C4", "accent2": "ED7D31", "accent3": "A5A5A5",
    "accent4": "FFC000", "accent5": "5B9BD5", "accent6": "70AD47",
    "hlink": "0563C1", "folHlink": "954F72",
}


def theme_colors(wb) -> dict[str, str]:
    colors = dict(_DEFAULT_THEME)
    raw = getattr(wb, "loaded_theme", None)
    if not raw:
        return colors
    try:
        ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        root = ET.fromstring(raw)
        scheme = root.find(".//a:clrScheme", ns)
        for child in list(scheme):
            name = child.tag.split("}")[1]
            srgb = child.find("a:srgbClr", ns)
            sys_ = child.find("a:sysClr", ns)
            if srgb is not None:
                colors[name] = srgb.get("val")
            elif sys_ is not None and sys_.get("lastClr"):
                colors[name] = sys_.get("lastClr")
    except Exception:
        pass
    return colors


def _apply_tint(hex6: str, tint: float) -> str:
    r, g, b = (int(hex6[i:i + 2], 16) / 255 for i in (0, 2, 4))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    if tint < 0:
        l = l * (1 + tint)
    elif tint > 0:
        l = l * (1 - tint) + tint
    r, g, b = colorsys.hls_to_rgb(h, l, s)
    return "".join(f"{round(c * 255):02X}" for c in (r, g, b))


def resolve_color(color, theme: dict[str, str]) -> str | None:
    """openpyxl Color → 'RRGGBB' (알 수 없으면 None)"""
    if color is None:
        return None
    try:
        if color.type == "rgb" and isinstance(color.rgb, str):
            return color.rgb[-6:].upper()
        if color.type == "theme" and color.theme is not None:
            idx = int(color.theme)
            if 0 <= idx < len(_THEME_ORDER):
                base = theme.get(_THEME_ORDER[idx])
                if base:
                    return _apply_tint(base, float(color.tint or 0))
        if color.type == "indexed" and color.indexed is not None:
            idx = int(color.indexed)
            if 0 <= idx < len(COLOR_INDEX):
                return COLOR_INDEX[idx][-6:].upper()
    except Exception:
        return None
    return None


def cell_fill_hex(cell, theme) -> str | None:
    fill = cell.fill
    if fill is None or fill.fill_type in (None, "none"):
        return None
    return resolve_color(fill.fgColor, theme)


def classify_hex(hex6: str | None) -> str:
    """'gray' | 'green' | 'yellow' | 'white' | 'none' | 'other'"""
    if not hex6:
        return "none"
    r, g, b = (int(hex6[i:i + 2], 16) / 255 for i in (0, 2, 4))
    h, s, v = colorsys.rgb_to_hsv(r, g, b)
    if s < 0.05:
        if v >= 0.99:
            return "white"
        return "gray" if v >= 0.35 else "other"
    deg = h * 360
    if 38 <= deg <= 72:
        return "yellow"
    if 73 <= deg <= 170:
        return "green"
    return "other"


KIND_KO = {"gray": "회색(프로그램)", "green": "초록(수식)", "yellow": "노랑(사람)",
           "white": "흰색", "none": "색 없음", "other": "기타색"}


# ── 시트 머리글 읽기 ─────────────────────────────────────
def find_header_row(ws, key_norm: str = "studyid", max_scan: int = 15) -> int | None:
    """key_norm 과 같은 글자가 있는 첫 행 번호 (없으면 None)."""
    for row in ws.iter_rows(min_row=1, max_row=min(max_scan, ws.max_row)):
        for cell in row:
            if norm(cell.value) == key_norm:
                return cell.row
    return None


def header_columns(ws, header_row: int, theme) -> list[dict]:
    """머리글 행의 각 열 정보: col, letter, text, norm, hex, kind"""
    cols = []
    for cell in ws[header_row]:
        if cell.value is None or str(cell.value).strip() == "":
            continue
        hex6 = cell_fill_hex(cell, theme)
        cols.append({
            "col": cell.column,
            "letter": get_column_letter(cell.column),
            "text": str(cell.value).strip(),
            "norm": norm(cell.value),
            "hex": hex6,
            "kind": classify_hex(hex6),
        })
    return cols


def list_validation_options(ws, col: int, row: int) -> list[str] | None:
    """(row, col) 칸에 걸린 목록형 데이터 유효성(드롭다운) 선택지. 없으면 None."""
    target = f"{get_column_letter(col)}{row}"
    for dv in ws.data_validations.dataValidation:
        if dv.type != "list" or not dv.formula1:
            continue
        if target in dv.sqref:
            f1 = dv.formula1.strip()
            if f1.startswith('"') and f1.endswith('"'):
                return [x.strip() for x in f1[1:-1].split(",")]
            return [f"(범위 참조: {f1})"]
    return None


def extend_tables(ws, last_row: int) -> None:
    """시트가 '표(Table)' 서식이면 새로 쓴 행까지 표 범위를 늘려 줍니다."""
    for tbl in ws.tables.values():
        min_col, min_row, max_col, max_row = range_boundaries(tbl.ref)
        if max_row < last_row:
            tbl.ref = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{last_row}"
            if tbl.autoFilter is not None:
                tbl.autoFilter.ref = tbl.ref


# ── 열기·백업·저장 ───────────────────────────────────────
def lock_file_exists(path: Path = config.EXCEL_PATH) -> bool:
    """엑셀이 파일을 열면 같은 폴더에 '~$파일명' 잠금 파일이 생깁니다."""
    names = {"~$" + path.name, "~$" + path.name[2:]}
    return any((path.parent / n).exists() for n in names)


def load(path: Path = config.EXCEL_PATH):
    """수식을 보존한 채로 엑셀을 읽습니다. (경고 목록도 함께 반환)"""
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        wb = openpyxl.load_workbook(path)
    return wb, [str(w.message) for w in caught]


def backup(path: Path = config.EXCEL_PATH, tag: str = "backup") -> Path:
    config.BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    dst = config.BACKUP_DIR / f"{path.stem}_{tag}_{stamp}{path.suffix}"
    shutil.copy2(path, dst)
    return dst


def save(wb, path: Path = config.EXCEL_PATH, retries: int = 3) -> None:
    """
    안전 저장: 임시 파일에 저장 → 원본과 교체.
    엑셀이 열려 있거나 OneDrive 동기화로 잠겨 있으면 몇 번 재시도 후 ExcelLockedError.
    """
    tmp = path.with_name(path.stem + ".saving.tmp")
    try:
        wb.save(tmp)
    except PermissionError as e:
        raise ExcelLockedError(f"임시 파일을 쓰지 못했습니다: {tmp}\n(상세: {e})") from e
    last_error = None
    for attempt in range(retries):
        try:
            os.replace(tmp, path)
            return
        except PermissionError as e:
            last_error = e
            time.sleep(0.7 * (attempt + 1))
    try:
        tmp.unlink()
    except OSError:
        pass
    raise ExcelLockedError(
        f"엑셀 파일을 저장하지 못했습니다.\n{path}\n\n"
        "Excel에서 이 파일이 열려 있으면 닫은 뒤 다시 시도하세요.\n"
        f"(상세: {last_error})"
    )
