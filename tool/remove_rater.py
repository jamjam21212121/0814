# -*- coding: utf-8 -*-
r"""
엑셀 양식에서 'rater' 열을 지웁니다. (한 번만 실행하면 되는 정리 도구)

안전장치
- 지우기 전에 labeling\backup\ 에 백업을 만듭니다.
- 다른 칸의 수식·드롭다운·조건부 서식이 지울 열이나 그 오른쪽 열을 가리키면
  (지우면 참조가 어긋나므로) 아무것도 바꾸지 않고 멈춥니다.
- 머리글이 아닌 칸에 'rater' 글자가 있으면 위치만 알려 줍니다(직접 확인용).

실행: fix_remove_rater.bat 더블클릭  (실행 전에 엑셀을 닫아 두세요)
"""
from __future__ import annotations

import sys

from openpyxl.formula.tokenizer import Token, Tokenizer
from openpyxl.utils import get_column_letter, range_boundaries

import config
import excel_io


def fail(msg: str) -> None:
    print("\n[중단] " + msg)
    sys.exit(1)


def formula_refs(formula: str, host: str):
    """수식 안의 (시트 이름, 범위) 목록."""
    try:
        tokens = Tokenizer(formula).items
    except Exception:
        return
    for t in tokens:
        if t.type == Token.OPERAND and t.subtype == Token.RANGE:
            value = t.value
            if "!" in value:
                sheet, rng = value.rsplit("!", 1)
                sheet = sheet.strip("'")
            else:
                sheet, rng = host, value
            yield sheet, rng.replace("$", "")


def max_col_of(rng: str) -> int | None:
    try:
        _, _, max_col, _ = range_boundaries(rng)
        return max_col
    except Exception:
        return None


def conflicts(wb, ws, col: int) -> list[str]:
    """ws 의 col 열 이상을 가리키는 수식/유효성/조건부서식."""
    found = []
    for other in wb.worksheets:
        for row in other.iter_rows():
            for cell in row:
                if not excel_io.is_formula(cell.value) or not isinstance(cell.value, str):
                    continue
                for sheet, rng in formula_refs(cell.value, other.title):
                    mc = max_col_of(rng)
                    if sheet == ws.title and mc is not None and mc >= col:
                        found.append(f"{other.title}!{cell.coordinate} 수식 {cell.value}")
    for dv in ws.data_validations.dataValidation:
        for rng in str(dv.sqref).split():
            mc = max_col_of(rng)
            if mc is not None and mc >= col:
                found.append(f"{ws.title} 데이터유효성 {rng}")
    for cf in ws.conditional_formatting:
        for rng in str(cf.sqref).split():
            mc = max_col_of(rng)
            if mc is not None and mc >= col:
                found.append(f"{ws.title} 조건부서식 {rng}")
    for mr in ws.merged_cells.ranges:
        if mr.max_col >= col:
            found.append(f"{ws.title} 병합셀 {mr}")
    for tbl in ws.tables.values():
        if max_col_of(tbl.ref) >= col:
            found.append(f"{ws.title} 표 {tbl.name} {tbl.ref}")
    return found


def delete_column(ws, col: int) -> None:
    """열 삭제 + 오른쪽 열 너비를 한 칸씩 당겨 맞춤."""
    last = ws.max_column
    widths = {c: ws.column_dimensions[get_column_letter(c)].width
              for c in range(col + 1, last + 1)
              if get_column_letter(c) in ws.column_dimensions}
    ws.delete_cols(col)
    for c in range(col, last + 1):
        letter = get_column_letter(c)
        if c + 1 in widths and widths[c + 1]:
            ws.column_dimensions[letter].width = widths[c + 1]
        elif letter in ws.column_dimensions:
            del ws.column_dimensions[letter]


def main() -> None:
    print("엑셀 양식에서 rater 열 지우기\n")
    if not config.EXCEL_PATH.exists():
        fail(f"엑셀 파일이 없습니다: {config.EXCEL_PATH}")
    if excel_io.lock_file_exists():
        print("[주의] Excel에서 파일이 열려 있는 것 같습니다. 닫지 않으면 저장이 실패합니다.")
    wb, _ = excel_io.load()

    targets = []   # (시트, 열 번호, 머리글 행)
    for ws in wb.worksheets:
        hr = excel_io.find_header_row(ws)
        if hr is None:
            continue
        for cell in ws[hr]:
            if excel_io.norm(cell.value) == "rater":
                targets.append((ws, cell.column, hr))
    mentions = [f"{ws.title}!{c.coordinate}: {c.value}"
                for ws in wb.worksheets for row in ws.iter_rows() for c in row
                if isinstance(c.value, str) and "rater" in c.value.lower()
                and not any(ws is t[0] and c.row == t[2] and c.column == t[1] for t in targets)]

    if not targets:
        print("rater 열이 없습니다. (이미 지웠거나 원래 없음) 바꿀 것이 없습니다.")
    else:
        for ws, col, hr in targets:
            print(f"- [{ws.title}] {get_column_letter(col)}열 '{ws.cell(hr, col).value}'")
        problems = [p for ws, col, _ in targets for p in conflicts(wb, ws, col)]
        if problems:
            fail("아래 항목이 지울 열(또는 그 오른쪽 열)을 가리키고 있어 자동으로 지우지 않았습니다.\n"
                 "        Excel에서 열 머리글을 오른쪽 클릭 → [삭제]로 직접 지워 주세요.\n          "
                 + "\n          ".join(problems[:20]))
        bk = excel_io.backup(tag="before_remove_rater")
        print(f"\n백업 저장: {bk}")
        # 같은 시트에 여러 개면 오른쪽부터 지워야 번호가 안 꼬임
        for ws, col, _ in sorted(targets, key=lambda t: -t[1]):
            delete_column(ws, col)
        try:
            excel_io.save(wb)
        except excel_io.ExcelLockedError as e:
            fail(str(e))
        print(f"rater 열 {len(targets)}개를 지웠습니다: {config.EXCEL_PATH}")

    if mentions:
        print("\n[참고] 머리글 말고 다른 칸에도 'rater' 글자가 있습니다. 필요하면 직접 고쳐 주세요:")
        for m in mentions:
            print("  " + m)


if __name__ == "__main__":
    main()
