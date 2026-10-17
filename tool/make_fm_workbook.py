# -*- coding: utf-8 -*-
"""
압박골절·전이암 라벨링 엑셀 양식 만들기 (한 번만 실행).

  labeling\\LumbarDISC_Fracture_Metastasis_labeling_v0.1.xlsx

- sample_list 는 쉬모를 엑셀의 sample_list 에서 같은 100명(순번·study_id·series_id)을 복사합니다.
  status 는 이 작업 기준으로 새로 '미완료' 부터 시작합니다.
- 이미 파일이 있으면 덮어쓰지 않고 멈춥니다 (기록 보호).
- 머리글 색 규칙은 쉬모를 양식과 같습니다:
  회색 = 프로그램이 기록 / 초록 = 자동 계산(수식) / 노란색 = 사람이 직접 입력
"""
from __future__ import annotations

import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import config
import excel_io
import tasks
import workbook

GRAY, GREEN, YELLOW = "D9D9D9", "C6EFCE", "FFFF00"
N_LABEL_ROWS = 101          # labels: 머리글 + 100명
N_LESION_ROWS = 301         # lesions: 머리글 + 300개

FR, MT = tasks.FM.condition("fracture"), tasks.FM.condition("metastasis")


def fill(hex6):
    return PatternFill("solid", fgColor="FF" + hex6)


def header(ws, items, freeze="B2"):
    """items: [(머리글, 색, 너비)]"""
    for i, (text, color, width) in enumerate(items, start=1):
        c = ws.cell(1, i, text)
        c.fill = fill(color)
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = freeze


def dropdown(ws, options, rng):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(rng)


def write_readme(wb):
    ws = wb.active
    ws.title = "README"
    lines = [
        "LumbarDISC 압박골절 · 전이암 라벨링 기록 양식 (초안 v0.1)",
        "",
        "대상: RSNA 2024 Lumbar Spine, 쉬모를 결절 연구와 같은 100명 (sample_list)",
        "판독 단위: 척추체 L1, L2, L3, L4, L5, S1 (상·하종판 구분 없음)",
        "기록 프로그램: tool\\run_fracture.bat  (쉬모를은 run_labeler.bat — 엑셀이 따로입니다)",
        "",
        "머리글 색 규칙",
        "  회색 = 프로그램이 기록하는 칸 (손으로 고치지 마세요)",
        "  초록 = 자동 계산(수식) 칸",
        "  노란색 = 사람이 직접 입력하는 칸 (review_status, final_label, note)",
        "",
        "시트",
        "  sample_list : 100명 목록과 이 작업의 진행 상태 (미완료 / 완료 / 애매포함)",
        "  labels      : 환자 1명 = 1행. fracture_L1~S1, metastasis_L1~S1 각 칸에 Normal / Fracture(또는 Metastasis) / Uncertain",
        "  lesions     : 소견 1개 = 1행. 어느 척추체에 무엇이 있는지 + 대표 좌표(x, y, instance_number)",
        "  summary     : 진행 현황과 소견 개수 (자동 계산)",
        "",
        "규칙",
        "  - 같은 척추체에 골절과 전이암이 함께 있으면 lesions 에 2행(condition 이 다름)으로 기록됩니다.",
        "  - 확실한 소견과 애매한 소견이 같은 척추체에 있으면 labels 칸에는 확실한 값이 들어갑니다.",
        "  - 애매(Uncertain) 소견은 review_status 가 자동으로 Pending 이 되고, 캡처가 labeling\\captures_fm\\ 에 저장됩니다.",
        "  - 아무 소견이 없으면 N 한 번으로 12칸이 모두 Normal 로 기록됩니다.",
    ]
    for i, t in enumerate(lines, start=1):
        ws.cell(i, 1, t)
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 110


def write_codebook(wb):
    ws = wb.create_sheet("codebook")
    rows = [
        ("칸", "값", "뜻"),
        ("labels fracture_*", "Normal", "압박골절 없음"),
        ("", "Fracture", "압박골절 있음 (확실)"),
        ("", "Uncertain", "애매 — 확인 필요"),
        ("labels metastasis_*", "Normal", "전이암 소견 없음"),
        ("", "Metastasis", "전이암 소견 있음 (확실)"),
        ("", "Uncertain", "애매 — 확인 필요"),
        ("lesions condition", FR.excel_name, "압박골절 (프로그램 키 F)"),
        ("", MT.excel_name, "전이암 (프로그램 키 C)"),
        ("lesions label", "Fracture / Metastasis", "확실"),
        ("", "Uncertain", "애매 (키 A) → review_status = Pending"),
        ("lesions review_status", "Pending", "확인 기다림 (자동)"),
        ("", "Confirmed", "확인 끝 (사람이 입력)"),
        ("lesions final_label", "Fracture / Metastasis / Normal", "확인 후 최종 판정 (사람이 입력)"),
        ("x, y", "픽셀", "원본 DICOM 픽셀 좌표 (가로 x, 세로 y), instance_number 는 슬라이스 번호"),
        ("series_id", "", "기록한 영상의 시리즈 (T2 또는 T1 보기에서 찍었으면 T1 시리즈)"),
    ]
    for r, row in enumerate(rows, start=1):
        for c, v in enumerate(row, start=1):
            ws.cell(r, c, v)
    for c in range(1, 4):
        ws.cell(1, c).font = Font(bold=True)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 32
    ws.column_dimensions["C"].width = 70


def write_sample_list(wb, patients):
    ws = wb.create_sheet("sample_list")
    header(ws, [("sample_no", GRAY, 11), ("study_id", GRAY, 14), ("sag_t2_series_id", GRAY, 18),
                ("sag_t1_series_id", GRAY, 18), ("status", GRAY, 11), ("labeled_at", GRAY, 18)])
    for i, p in enumerate(patients, start=2):
        ws.cell(i, 1, workbook._num(p.sample_no) if p.sample_no else i - 1)
        ws.cell(i, 2, workbook._num(p.study_id))
        ws.cell(i, 3, workbook._num(p.t2_series) if p.t2_series else None)
        ws.cell(i, 4, workbook._num(p.t1_series) if p.t1_series else None)
        ws.cell(i, 5, config.STATUS_TODO)
    last = max(len(patients) + 1, 2)
    dropdown(ws, [config.STATUS_TODO, config.STATUS_DONE, config.STATUS_UNCERTAIN], f"E2:E{last}")


def write_labels(wb):
    ws = wb.create_sheet("labels")
    items = [("study_id", GRAY, 13)]
    items += [(f"fracture_{v}", GRAY, 12) for v in tasks.VERTEBRAE]          # B~G
    items += [(f"metastasis_{v}", GRAY, 13) for v in tasks.VERTEBRAE]        # H~M
    items += [("n_fracture", GREEN, 11), ("n_metastasis", GREEN, 12), ("n_uncertain", GREEN, 11),   # N O P
              ("patient_fracture", GREEN, 15), ("patient_metastasis", GREEN, 17),                     # Q R
              ("labeled_at", GRAY, 17), ("note", YELLOW, 30)]                                         # S T
    header(ws, items)
    for r in range(2, N_LABEL_ROWS + 1):
        ws[f"N{r}"] = f'=IF($A{r}="","",COUNTIF(B{r}:G{r},"Fracture"))'
        ws[f"O{r}"] = f'=IF($A{r}="","",COUNTIF(H{r}:M{r},"Metastasis"))'
        ws[f"P{r}"] = f'=IF($A{r}="","",COUNTIF(B{r}:M{r},"Uncertain"))'
        ws[f"Q{r}"] = (f'=IF($A{r}="","",IF(N{r}>0,"Fracture",'
                       f'IF(COUNTIF(B{r}:G{r},"Uncertain")>0,"Uncertain","Normal")))')
        ws[f"R{r}"] = (f'=IF($A{r}="","",IF(O{r}>0,"Metastasis",'
                       f'IF(COUNTIF(H{r}:M{r},"Uncertain")>0,"Uncertain","Normal")))')
    dropdown(ws, ["Normal", "Fracture", "Uncertain"], f"B2:G{N_LABEL_ROWS}")
    dropdown(ws, ["Normal", "Metastasis", "Uncertain"], f"H2:M{N_LABEL_ROWS}")


def write_lesions(wb):
    ws = wb.create_sheet("lesions")
    header(ws, [("lesion_id", GRAY, 14), ("study_id", GRAY, 12), ("series_id", GRAY, 12),     # A B C
                ("instance_number", GRAY, 15), ("condition", GRAY, 21), ("vertebra", GRAY, 9),  # D E F
                ("x", GRAY, 8), ("y", GRAY, 8), ("label", GRAY, 11), ("capture_file", GRAY, 34),  # G H I J
                ("review_status", YELLOW, 13), ("final_label", YELLOW, 12), ("note", YELLOW, 30),   # K L M
                ("labeled_at", GRAY, 17)])                                                          # N
    n = N_LESION_ROWS
    dropdown(ws, [FR.excel_name, MT.excel_name], f"E2:E{n}")
    dropdown(ws, tasks.VERTEBRAE, f"F2:F{n}")
    dropdown(ws, ["Fracture", "Metastasis", "Uncertain"], f"I2:I{n}")
    dropdown(ws, ["Pending", "Confirmed"], f"K2:K{n}")
    dropdown(ws, ["Fracture", "Metastasis", "Normal"], f"L2:L{n}")


def write_summary(wb):
    ws = wb.create_sheet("summary")
    L = f"labels!$A$2:$A${N_LABEL_ROWS}"
    rows = [
        ("압박골절 · 전이암 진행 현황 (자동 계산)", None),
        (None, None),
        ("판독한 환자 수", f"=COUNTA({L})"),
        ("미완료 (sample_list)", '=COUNTIF(sample_list!$E:$E,"미완료")'),
        ("확인필요 (sample_list 애매포함)", '=COUNTIF(sample_list!$E:$E,"애매포함")'),
        (None, None),
        ("압박골절 있는 환자", f'=COUNTIF(labels!$Q$2:$Q${N_LABEL_ROWS},"Fracture")'),
        ("압박골절 애매 환자", f'=COUNTIF(labels!$Q$2:$Q${N_LABEL_ROWS},"Uncertain")'),
        ("압박골절 척추체 수", f'=SUM(labels!$N$2:$N${N_LABEL_ROWS})'),
        ("전이암 있는 환자", f'=COUNTIF(labels!$R$2:$R${N_LABEL_ROWS},"Metastasis")'),
        ("전이암 애매 환자", f'=COUNTIF(labels!$R$2:$R${N_LABEL_ROWS},"Uncertain")'),
        ("전이암 척추체 수", f'=SUM(labels!$O$2:$O${N_LABEL_ROWS})'),
        (None, None),
        ("확인 대기 (Pending)", f'=COUNTIF(lesions!$K$2:$K${N_LESION_ROWS},"Pending")'),
        (None, None),
        ("척추체별", "압박골절"),
    ]
    for r, (a, b) in enumerate(rows, start=1):
        ws.cell(r, 1, a)
        ws.cell(r, 2, b)
    ws.cell(len(rows), 3, "전이암")
    for i, v in enumerate(tasks.VERTEBRAE):
        r = len(rows) + 1 + i
        fc, mc = get_column_letter(2 + i), get_column_letter(8 + i)
        ws.cell(r, 1, v)
        ws.cell(r, 2, f'=COUNTIF(labels!${fc}$2:${fc}${N_LABEL_ROWS},"Fracture")')
        ws.cell(r, 3, f'=COUNTIF(labels!${mc}$2:${mc}${N_LABEL_ROWS},"Metastasis")')
    ws["A1"].font = Font(bold=True, size=13)
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 12


def build(patients, path):
    wb = Workbook()
    write_readme(wb)
    write_codebook(wb)
    write_sample_list(wb, patients)
    write_labels(wb)
    write_lesions(wb)
    write_summary(wb)
    path.parent.mkdir(parents=True, exist_ok=True)
    excel_io.save(wb, path)


def main() -> int:
    print(f"코드 버전: {config.TOOL_VERSION}")
    dst = config.FM_EXCEL_PATH
    if dst.exists():
        print(f"이미 있습니다 — 기록 보호를 위해 덮어쓰지 않습니다:\n  {dst}")
        print("새로 만들려면 그 파일을 다른 곳으로 옮긴 뒤 다시 실행하세요.")
        return 1
    if not config.EXCEL_PATH.exists():
        print(f"쉬모를 엑셀이 없습니다 (sample_list 를 복사해 와야 합니다):\n  {config.EXCEL_PATH}")
        return 1
    try:
        patients = workbook.read_sample_list(config.EXCEL_PATH)
    except Exception as e:
        print(f"쉬모를 엑셀 sample_list 를 읽지 못했습니다: {e}")
        return 1
    if not patients:
        print("쉬모를 엑셀 sample_list 가 비어 있습니다. 먼저 step1_make_sample.bat 을 실행하세요.")
        return 1
    try:
        build(patients, dst)
    except excel_io.ExcelLockedError as e:
        print(e)
        return 1
    print(f"만들었습니다 ({len(patients)}명): {dst}")
    print("이제 run_fracture.bat 으로 판독을 시작하면 됩니다.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
